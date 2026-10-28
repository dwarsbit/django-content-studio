"""
Export: a studio's list as a CSV or Excel (.xlsx) file.

A model studio opts in with export_fields (the columns, like
table_display) or allow_export (which exports the table_display
columns). Cells hold the values as the admin displays them: choice
labels, the string of related objects, dates in ISO format. CSV is
streamed and always available; Excel needs openpyxl (the "excel" extra)
and is built in memory.
"""

import csv
import datetime
import html
import importlib.util
import json
import uuid
from decimal import Decimal
from io import BytesIO

from django.core.exceptions import FieldDoesNotExist
from django.core.serializers.json import DjangoJSONEncoder
from django.db import models
from django.db.models.fields.files import FieldFile
from django.http import HttpResponse, StreamingHttpResponse
from django.utils import timezone
from django.utils.functional import Promise
from django.utils.html import strip_tags
from django.utils.http import content_disposition_header
from django.utils.text import capfirst, slugify
from django.utils.translation import gettext as _

from .fields import HTMLField, MultipleChoiceField
from .form import TableColumn

CSV = "csv"
XLSX = "xlsx"

CONTENT_TYPES = {
    CSV: "text/csv; charset=utf-8",
    XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

# Cells starting with one of these are taken as a formula by spreadsheet
# applications (CSV injection); such cells get a leading apostrophe.
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

# Rows fetched per database round trip while streaming.
CHUNK_SIZE = 500


def is_excel_available() -> bool:
    """Whether Excel export is possible: openpyxl is installed."""
    return importlib.util.find_spec("openpyxl") is not None


def get_export_formats() -> list[str]:
    """The formats exports are offered in, CSV first."""
    return [CSV, XLSX] if is_excel_available() else [CSV]


def is_exportable(studio) -> bool:
    """
    Whether a studio offers export: a model studio that sets
    export_fields or allow_export and is not a singleton. Inlines are
    never exported on their own.
    """
    from .admin import ModelStudio, get_is_singleton

    return (
        isinstance(studio, ModelStudio)
        and bool(getattr(studio, "export_fields", None) or studio.allow_export)
        and not get_is_singleton(studio)
    )


def get_export_entries(studio) -> list:
    """The studio's export columns as configured: export_fields, or the
    table_display columns with allow_export."""
    from .admin import resolve_table_display

    export_fields = getattr(studio, "export_fields", None)

    if export_fields:
        return list(export_fields)

    return resolve_table_display(studio)


def get_export_columns(studio) -> list[TableColumn]:
    """
    The columns an export writes, as TableColumn objects. Fields in the
    studio's exclude are left out: they never leave the server.
    """
    excluded = set(studio.exclude or [])
    columns = [
        entry if isinstance(entry, TableColumn) else TableColumn(field=entry)
        for entry in get_export_entries(studio)
    ]

    return [column for column in columns if column.field not in excluded]


def get_model_field(studio, column: TableColumn):
    """The forward model field a column names, or None (for __str__,
    studio methods and reverse relations)."""
    from .admin import is_computed_column

    if column.field == "__str__" or is_computed_column(studio, column):
        return None

    try:
        field = studio.model._meta.get_field(column.field)
    except FieldDoesNotExist:
        return None

    if field.auto_created and not field.concrete:
        return None

    return field


def get_header(studio, column: TableColumn) -> str:
    """A column's header: its explicit header, the field's capitalized
    verbose name, or the model's for __str__."""
    if column.header:
        return str(column.header)

    model = studio.model

    if column.field == "__str__":
        return capfirst(str(model._meta.verbose_name))

    field = get_model_field(studio, column)

    if field is not None:
        return capfirst(str(field.verbose_name))

    return capfirst(column.field.replace("_", " ").strip())


def get_field_value(obj, field):
    """A model field's value as the admin displays it."""
    if field.many_to_many:
        return ", ".join(str(related) for related in getattr(obj, field.name).all())

    if field.is_relation:
        related = getattr(obj, field.name)
        return None if related is None else str(related)

    value = field.value_from_object(obj)

    if field.flatchoices and value is not None and not isinstance(value, list):
        return getattr(obj, f"get_{field.name}_display")()

    if isinstance(field, MultipleChoiceField) and isinstance(value, list):
        labels = {option: label for option, label in field.options}
        return ", ".join(str(labels.get(option, option)) for option in value)

    if isinstance(field, HTMLField) and isinstance(value, str):
        return html.unescape(strip_tags(value)).strip()

    return value


def get_value(studio, column: TableColumn, obj, request):
    """The raw value of a column for one object."""
    if column.field == "__str__":
        return str(obj)

    field = get_model_field(studio, column)

    if field is not None:
        return get_field_value(obj, field)

    return getattr(studio, column.field)(obj, request)


def to_cell(value, native: bool = False):
    """
    A value as a cell: text, or with native (Excel) numbers and dates
    as they are. Booleans read Yes or No, aware datetimes are in the
    current time zone, lists of plain values are joined with commas and
    other structures written as JSON. Text that a spreadsheet would read
    as a formula is prefixed with an apostrophe.
    """
    if value is None:
        return ""

    if isinstance(value, bool):
        value = _("Yes") if value else _("No")
    elif isinstance(value, (int, float, Decimal)):
        return value if native else str(value)
    elif isinstance(value, datetime.datetime):
        if timezone.is_aware(value):
            value = timezone.localtime(value)

        if native:
            return value.replace(tzinfo=None)

        return value.isoformat()
    elif isinstance(value, (datetime.date, datetime.time)):
        return value if native else value.isoformat()
    elif isinstance(value, datetime.timedelta):
        value = str(value)
    elif isinstance(value, FieldFile):
        value = value.name or ""
    elif isinstance(value, (list, tuple)) and all(
        isinstance(item, (str, int, float, Decimal, Promise)) for item in value
    ):
        value = ", ".join(str(item) for item in value)
    elif isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, cls=DjangoJSONEncoder, ensure_ascii=False)
    elif isinstance(value, (uuid.UUID, Promise)):
        value = str(value)
    elif isinstance(value, models.Model):
        value = str(value)

    value = str(value)

    if value.startswith(FORMULA_PREFIXES):
        value = f"'{value}"

    return value


def prepare_queryset(studio, queryset):
    """Fetch the related objects the columns show along with the rows."""
    select, prefetch = [], []

    for column in get_export_columns(studio):
        field = get_model_field(studio, column)

        if field is None or not field.is_relation:
            continue

        if field.many_to_many:
            prefetch.append(field.name)
        else:
            select.append(field.name)

    if select:
        queryset = queryset.select_related(*select)

    if prefetch:
        queryset = queryset.prefetch_related(*prefetch)

    return queryset


def iter_rows(studio, request, queryset, native: bool = False):
    """The header row, then one row per object."""
    columns = get_export_columns(studio)

    yield [get_header(studio, column) for column in columns]

    for obj in prepare_queryset(studio, queryset).iterator(chunk_size=CHUNK_SIZE):
        yield [
            to_cell(get_value(studio, column, obj, request), native=native)
            for column in columns
        ]


def get_filename(studio, extension: str) -> str:
    """The download's name: the model's plural name and today's date."""
    name = slugify(str(studio.model._meta.verbose_name_plural)) or "export"
    return f"{name}-{timezone.localdate().isoformat()}.{extension}"


class _Echo:
    """A file-like object that returns what is written, for streaming
    the csv writer's lines."""

    def write(self, value):
        return value


def csv_response(studio, request, queryset) -> StreamingHttpResponse:
    """
    The export as a streamed CSV file, UTF-8 with a byte order mark so
    spreadsheet applications detect the encoding.
    """
    writer = csv.writer(_Echo())

    def stream():
        yield "﻿"

        for row in iter_rows(studio, request, queryset):
            yield writer.writerow(row)

    response = StreamingHttpResponse(stream(), content_type=CONTENT_TYPES[CSV])
    response["Content-Disposition"] = content_disposition_header(
        True, get_filename(studio, CSV)
    )
    return response


def xlsx_response(studio, request, queryset) -> HttpResponse:
    """
    The export as an Excel workbook with one sheet. Numbers and dates are
    native cells; dates show in ISO format.
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font

    workbook = Workbook(write_only=True)
    # Sheet titles are at most 31 characters, without []:*?/\.
    title = "".join(
        char
        for char in capfirst(str(studio.model._meta.verbose_name_plural))
        if char not in "[]:*?/\\"
    )
    sheet = workbook.create_sheet(title=title[:31] or "Export")
    bold = Font(bold=True)

    def cell(value, header=False):
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)

        result = WriteOnlyCell(sheet, value=value)

        if header:
            result.font = bold
        elif isinstance(value, datetime.datetime):
            result.number_format = "yyyy-mm-dd hh:mm:ss"
        elif isinstance(value, datetime.date):
            result.number_format = "yyyy-mm-dd"
        elif isinstance(value, datetime.time):
            result.number_format = "hh:mm:ss"
        elif isinstance(value, str):
            # Text stays text: never a formula.
            result.data_type = "s"

        return result

    rows = iter_rows(studio, request, queryset, native=True)
    sheet.append([cell(header, header=True) for header in next(rows)])

    for row in rows:
        sheet.append([cell(value) for value in row])

    output = BytesIO()
    workbook.save(output)

    response = HttpResponse(output.getvalue(), content_type=CONTENT_TYPES[XLSX])
    response["Content-Disposition"] = content_disposition_header(
        True, get_filename(studio, XLSX)
    )
    return response


def export_response(studio, request, queryset, export_format: str):
    """The export of queryset in the given format, as a download."""
    if export_format == XLSX:
        return xlsx_response(studio, request, queryset)

    return csv_response(studio, request, queryset)
